"""
Jev Test Runner
---------------
Runs manual test cases written in Excel (one row per step) using:
  - Playwright  -> drives the browser
  - TypeSafe Jev -> decides the action, the target element and the test data value
                    for each step, and checks the Expected Result
Then writes PASS / FAIL / REVIEW / SKIPPED / ERROR back into a Status column for every step.

Usage:
    Put your key in a .env file next to this script:
        TYPESAFE_API_KEY=your-key
    python JevTestRunner.py manual_test_cases.xlsx
"""

import os
import re
import sys
import time
from datetime import datetime
from pathlib import Path

from dotenv import load_dotenv
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from playwright.sync_api import sync_playwright
from typesafe_sdk import Choice, Noul, TypeSafeClient

# ----------------------------------------------------------------------------
# Settings
# ----------------------------------------------------------------------------
load_dotenv(Path(__file__).with_name(".env"))
TYPESAFE_API_KEY = os.getenv("TYPESAFE_API_KEY", "").strip()

INPUT_FILE = sys.argv[1] if len(sys.argv) > 1 else "manual_test_cases.xlsx"
SHEET_NAME = None            # None = first sheet
HEADLESS = False             # keep False for the video
BROWSER_CHANNEL = "chrome"    # use installed Google Chrome ("msedge" for Edge, None for Playwright's own Chromium)
SLOW_MO_MS = 0               # 0 = full speed; raise (e.g. 400) to slow actions so viewers can follow

DECISION_MIN_CONFIDENCE = 0.60   # below this, Jev's action/target pick is not executed
PASS_PROBABILITY = 0.80          # expected-result probability at or above this -> PASS
FAIL_PROBABILITY = 0.20          # at or below this -> FAIL, in between -> REVIEW

MAX_ELEMENTS = 80            # interactive elements sent to Jev per step
MAX_PAGE_TEXT = 5000         # characters of visible page text sent for verification
PRICE_PER_BILLION_INPUT = 42.0   # TypeSafe's published input price (USD); output is free

# Excel header names (matched case-insensitively). Rename here if your sheet differs.
COLUMNS = {
    "tc_id": "TC ID",
    "title": "Test Case Title",
    "url": "Start URL",
    "step_no": "Step No",
    "step": "Step Description",
    "data": "Test Data",
    "expected": "Expected Result",
    "status": "Status",
    "confidence": "Confidence",
    "remarks": "Remarks",
}

STATUS_FILLS = {
    "PASS": "C6EFCE",
    "FAIL": "FFC7CE",
    "ERROR": "FFC7CE",
    "REVIEW": "FFEB9C",
    "SKIPPED": "D9D9D9",
}

ACTIONS = {
    "click": "Click a button, link, icon, tab or checkbox",
    "fill": "Type text into an input box or text area",
    "select": "Choose an option from a dropdown list",
    "press_enter": "Press the Enter key in a field",
    "verify_only": "No interaction; the step only checks what is on screen",
}

# ----------------------------------------------------------------------------
# Browser helpers (JavaScript run inside the page)
# ----------------------------------------------------------------------------
COLLECT_ELEMENTS_JS = """
(maxCount) => {
  document.querySelectorAll('[data-jev-id]').forEach(e => e.removeAttribute('data-jev-id'));
  const sel = 'a, button, input, select, textarea, [role="button"], [role="link"], ' +
              '[role="checkbox"], [role="tab"], [role="menuitem"], [onclick]';
  const out = [];
  let i = 0;
  for (const el of document.querySelectorAll(sel)) {
    const r = el.getBoundingClientRect();
    const st = getComputedStyle(el);
    if (r.width === 0 || r.height === 0 || st.visibility === 'hidden' || st.display === 'none') continue;
    if ((el.getAttribute('type') || '') === 'hidden') continue;
    const id = 'e' + (i++);
    el.setAttribute('data-jev-id', id);
    let label = el.getAttribute('aria-label') || el.getAttribute('placeholder') || '';
    if (!label && el.id) {
      const l = document.querySelector('label[for="' + CSS.escape(el.id) + '"]');
      if (l) label = l.innerText;
    }
    if (!label) label = (el.innerText || el.value || el.getAttribute('title') || '').trim();
    const card = el.closest('li, tr, article, [class*="item"], [class*="card"]');
    const ctx = card ? (card.innerText || '').split('\\n')[0].trim() : '';
    out.push({
      id,
      tag: el.tagName.toLowerCase(),
      type: el.getAttribute('type') || '',
      label: label.replace(/\\s+/g, ' ').slice(0, 80),
      elId: el.id || '',
      name: el.getAttribute('name') || '',
      cls: (typeof el.className === 'string' ? el.className : '').slice(0, 60),
      ctx: ctx.slice(0, 60),
    });
    if (out.length >= maxCount) break;
  }
  return out;
}
"""

PAGE_SNAPSHOT_JS = """
() => {
  const fields = [];
  document.querySelectorAll('input, textarea, select').forEach(el => {
    const t = (el.getAttribute('type') || '').toLowerCase();
    if (['hidden', 'submit', 'button'].includes(t)) return;
    const name = el.getAttribute('aria-label') || el.getAttribute('placeholder') ||
                 el.getAttribute('name') || el.id || el.tagName.toLowerCase();
    const value = t === 'password' ? (el.value ? '(filled, hidden)' : '(empty)') : (el.value || '(empty)');
    fields.push(name + ': ' + value);
  });
  return { text: document.body ? document.body.innerText : '', fields: fields.join('\\n') };
}
"""


def describe_element(el: dict) -> str:
    bits = [el["tag"]]
    if el["type"]:
        bits.append(f'type={el["type"]}')
    if el["label"]:
        bits.append(f'"{el["label"]}"')
    if el["elId"]:
        bits.append(f'id={el["elId"]}')
    if el["name"]:
        bits.append(f'name={el["name"]}')
    if el["cls"]:
        bits.append(f'class={el["cls"]}')
    if el["ctx"]:
        bits.append(f'(inside: {el["ctx"]})')
    return " ".join(bits)


def parse_test_data(raw) -> list[tuple[str, str]]:
    """'username=standard_user; password=secret' -> [('username','standard_user'), ...]
       'standard_user' -> [('', 'standard_user')]"""
    if raw is None or not str(raw).strip():
        return []
    pairs = []
    for part in re.split(r"[;\n]", str(raw)):
        part = part.strip()
        if not part:
            continue
        if "=" in part:
            k, v = part.split("=", 1)
            pairs.append((k.strip(), v.strip()))
        else:
            pairs.append(("", part))
    return pairs


# ----------------------------------------------------------------------------
# Excel helpers
# ----------------------------------------------------------------------------
def map_columns(ws) -> dict:
    header = {str(c.value).strip().lower(): c.column for c in ws[1] if c.value}
    cols = {}
    for key, name in COLUMNS.items():
        if name.lower() in header:
            cols[key] = header[name.lower()]
        elif key in ("status", "confidence", "remarks"):
            new_col = ws.max_column + 1
            ws.cell(row=1, column=new_col, value=name).font = Font(name="Arial", bold=True)
            cols[key] = new_col
        else:
            raise SystemExit(f'Missing required column "{name}" in row 1 of the sheet.')
    return cols


def load_test_cases(ws, cols) -> list[dict]:
    """Groups step rows by TC ID. TC ID / Title / Start URL can be left blank on
    later steps of the same test case; they are carried forward."""
    cases, current = [], None
    last = {"tc_id": None, "title": None, "url": None}
    for r in range(2, ws.max_row + 1):
        step = ws.cell(r, cols["step"]).value
        if step is None or not str(step).strip():
            continue
        for k in last:
            v = ws.cell(r, cols[k]).value
            if v is not None and str(v).strip():
                last[k] = str(v).strip()
        if current is None or current["tc_id"] != last["tc_id"]:
            current = {"tc_id": last["tc_id"], "title": last["title"], "url": last["url"], "steps": []}
            cases.append(current)
        current["steps"].append({
            "row": r,
            "no": ws.cell(r, cols["step_no"]).value,
            "step": str(step).strip(),
            "data": parse_test_data(ws.cell(r, cols["data"]).value),
            "expected": str(ws.cell(r, cols["expected"]).value or "").strip(),
        })
    return cases


def write_result(ws, cols, row, status, confidence, remarks):
    cell = ws.cell(row=row, column=cols["status"], value=status)
    cell.font = Font(name="Arial", bold=True)
    cell.fill = PatternFill("solid", start_color=STATUS_FILLS.get(status, "FFFFFF"))
    conf = ws.cell(row=row, column=cols["confidence"], value=round(confidence, 3) if confidence is not None else None)
    conf.number_format = "0.0%"
    conf.font = Font(name="Arial")
    ws.cell(row=row, column=cols["remarks"], value=remarks).font = Font(name="Arial")


def add_summary_sheet(wb, ws, cols):
    name = "Run Summary"
    if name in wb.sheetnames:
        del wb[name]
    s = wb.create_sheet(name)
    col_letter = ws.cell(row=1, column=cols["status"]).column_letter
    rng = f"'{ws.title}'!${col_letter}$2:${col_letter}${ws.max_row}"
    s["A1"], s["B1"] = "Status", "Steps"
    for i, status in enumerate(["PASS", "FAIL", "REVIEW", "SKIPPED", "ERROR"], start=2):
        s[f"A{i}"] = status
        s[f"B{i}"] = f'=COUNTIF({rng},"{status}")'
        s[f"A{i}"].fill = PatternFill("solid", start_color=STATUS_FILLS[status])
    s["A8"], s["B8"] = "Total", "=SUM(B2:B6)"
    s["A9"], s["B9"] = "Pass rate", "=IF(B8=0,0,B2/B8)"
    s["B9"].number_format = "0.0%"
    for row in s["A1:B9"]:
        for c in row:
            c.font = Font(name="Arial", bold=c.row in (1, 8, 9))
    s.column_dimensions["A"].width = 14


# ----------------------------------------------------------------------------
# Jev calls
# ----------------------------------------------------------------------------
class Usage:
    def __init__(self):
        self.calls = 0
        self.input_tokens = 0

    def add(self, resp):
        self.calls += 1
        if resp.usage and resp.usage.input_tokens:
            self.input_tokens += resp.usage.input_tokens


def decide_step(client, usage, page, step, elements):
    """One Jev call: which action, which element, which test-data value."""
    data = step["data"]
    state = {
        "manual_test_step": step["step"],
        "test_data": "; ".join(f"{k}={v}" if k else v for k, v in data) or "(none)",
        "page_url": page.url,
        "page_title": page.title(),
    }
    questions = {
        "action": Choice(
            instructions="Which UI action does this manual test step ask the tester to perform?",
            criteria=ACTIONS,
        ),
    }
    if elements:
        questions["target"] = Choice(
            instructions="Which on-screen element should this test step act on?",
            criteria={e["id"]: describe_element(e) for e in elements},
        )
    if len(data) > 1:
        questions["value"] = Choice(
            instructions="Which test data value should be typed or selected in this step?",
            criteria={f"v{i}": (f"{k} = {v}" if k else v) for i, (k, v) in enumerate(data)},
        )
    t0 = time.perf_counter()
    resp = client.system_one(state=state, questions=questions)
    usage.add(resp)
    return resp, (time.perf_counter() - t0) * 1000


def verify_expected(client, usage, page, step):
    """One Jev call: probability that the Expected Result is now true."""
    page.wait_for_load_state("domcontentloaded")
    page.wait_for_timeout(500)
    snap = page.evaluate(PAGE_SNAPSHOT_JS)
    state = {
        "page_url": page.url,
        "page_title": page.title(),
        "visible_text": snap["text"][:MAX_PAGE_TEXT],
        "form_field_values": snap["fields"] or "(none)",
    }
    # Ask about the page as a plain statement. Mentioning the step made Jev doubt whether
    # the step caused the change, which held true results around p=0.5-0.7.
    question = Noul(
        instructions=f'Is this statement about the current web page true? "{step["expected"]}"'
    )
    resp = client.system_one(state=state, questions={"expected_met": question})
    usage.add(resp)
    return resp.nouls["expected_met"].noul


# ----------------------------------------------------------------------------
# Step execution
# ----------------------------------------------------------------------------
def run_step(client, usage, page, step):
    """Returns (status, confidence, remarks, stop_test_case)."""
    elements = page.evaluate(COLLECT_ELEMENTS_JS, MAX_ELEMENTS)
    resp, ms = decide_step(client, usage, page, step, elements)

    action = resp.choices["action"]
    remarks = [f"{action.choice} ({action.confidence:.2f})"]
    decision_conf = action.confidence

    if action.choice != "verify_only":
        if "target" not in resp.choices:
            return "REVIEW", decision_conf, "No interactive elements found on the page", True
        target = resp.choices["target"]
        decision_conf = min(decision_conf, target.confidence)
        el = next(e for e in elements if e["id"] == target.choice)
        remarks.append(f'on {describe_element(el)[:70]} ({target.confidence:.2f})')

        if decision_conf < DECISION_MIN_CONFIDENCE:
            remarks.append("not executed: low confidence, needs human review")
            return "REVIEW", decision_conf, " | ".join(remarks), True

        value = None
        if action.choice in ("fill", "select"):
            if not step["data"]:
                return "REVIEW", decision_conf, "Step needs test data but the Test Data cell is empty", True
            if len(step["data"]) == 1:
                value = step["data"][0][1]
            else:
                v = resp.choices["value"]
                decision_conf = min(decision_conf, v.confidence)
                value = step["data"][int(v.choice[1:])][1]
                if v.confidence < DECISION_MIN_CONFIDENCE:
                    return "REVIEW", decision_conf, f"Unsure which test data value to use ({v.confidence:.2f})", True
            remarks.append(f'value "{value}"')

        loc = page.locator(f'[data-jev-id="{target.choice}"]')
        if action.choice == "click":
            loc.click()
        elif action.choice == "fill":
            loc.fill(value)
        elif action.choice == "select":
            loc.select_option(label=value)
        elif action.choice == "press_enter":
            loc.press("Enter")

    remarks.append(f"jev {ms:.0f} ms")

    if not step["expected"]:
        return "PASS", decision_conf, " | ".join(remarks + ["action executed, no expected result given"]), False

    p = verify_expected(client, usage, page, step)
    remarks.append(f"expected result p={p:.2f}")
    if p >= PASS_PROBABILITY:
        return "PASS", p, " | ".join(remarks), False
    if p <= FAIL_PROBABILITY:
        return "FAIL", p, " | ".join(remarks), True
    return "REVIEW", p, " | ".join(remarks + ["unclear, needs human review"]), False


# ----------------------------------------------------------------------------
# Main
# ----------------------------------------------------------------------------
def main():
    if not TYPESAFE_API_KEY or TYPESAFE_API_KEY == "your-api-key-here":
        sys.exit("TYPESAFE_API_KEY is not set. Add it to the .env file next to JevTestRunner.py.")

    src = Path(INPUT_FILE)
    wb = load_workbook(src)
    ws = wb[SHEET_NAME] if SHEET_NAME else wb.worksheets[0]
    cols = map_columns(ws)
    cases = load_test_cases(ws, cols)

    out_dir = Path("results")
    shots = out_dir / "screenshots"
    shots.mkdir(parents=True, exist_ok=True)
    out_file = out_dir / f"{src.stem}_results_{datetime.now():%Y%m%d_%H%M%S}.xlsx"

    usage = Usage()
    counts = {}
    print(f"\nRunning {len(cases)} test cases, {sum(len(c['steps']) for c in cases)} steps\n")

    with TypeSafeClient(api_key=TYPESAFE_API_KEY) as client, sync_playwright() as pw:
        browser = pw.chromium.launch(channel=BROWSER_CHANNEL, headless=HEADLESS, slow_mo=SLOW_MO_MS)
        for case in cases:
            print(f"== {case['tc_id']}  {case['title'] or ''}")
            context = browser.new_context(viewport={"width": 1280, "height": 800})
            page = context.new_page()
            page.set_default_timeout(10_000)
            stop = False
            start_error = None
            try:
                if case["url"]:
                    page.goto(case["url"])
            except Exception as exc:
                start_error = f"Could not open start URL: {str(exc)[:200]}"

            for i, step in enumerate(case["steps"]):
                if start_error:
                    status, conf = ("ERROR" if i == 0 else "SKIPPED"), None
                    remarks = start_error if i == 0 else "Skipped: start URL could not be opened"
                elif stop:
                    status, conf, remarks = "SKIPPED", None, "Skipped after an earlier step did not pass"
                else:
                    try:
                        status, conf, remarks, stop = run_step(client, usage, page, step)
                    except Exception as exc:
                        status, conf, remarks, stop = "ERROR", None, f"{type(exc).__name__}: {str(exc)[:200]}", True
                    if status in ("FAIL", "ERROR"):
                        shot = shots / f"{case['tc_id']}_step{step['no']}.png"
                        try:
                            page.screenshot(path=str(shot), full_page=True)
                            remarks += f" | screenshot: {shot}"
                        except Exception:
                            pass

                write_result(ws, cols, step["row"], status, conf, remarks)
                counts[status] = counts.get(status, 0) + 1
                conf_txt = f"{conf:.2f}" if conf is not None else "  - "
                print(f"   step {step['no']!s:<3} {status:<8} {conf_txt}  {remarks[:110]}")

            context.close()
            add_summary_sheet(wb, ws, cols)
            wb.save(out_file)   # save after every test case so progress is never lost
        browser.close()

    cost = usage.input_tokens / 1e9 * PRICE_PER_BILLION_INPUT
    print("\n" + "-" * 60)
    print("  ".join(f"{k}: {v}" for k, v in sorted(counts.items())))
    print(f"Jev calls: {usage.calls}   input tokens: {usage.input_tokens:,}   est. cost: ${cost:.6f}")
    print(f"Results written to: {out_file}\n")


if __name__ == "__main__":
    main()